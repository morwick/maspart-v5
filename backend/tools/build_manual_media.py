# -*- coding: utf-8 -*-
"""Ekstrak GAMBAR (diagram/skema/pinout/foto unit) dari sumber pengetahuan mentah
di `data/manuals/` → `data/manual_media/*.png` + `index.json`.

Selama ini sumber PDF/Excel hanya diambil TABEL-nya (pin_ecu, jadwal_perawatan,
filter_shantui, abs_scr). Gambar di dalamnya — skema kelistrikan ECU, pinout
konektor, foto unit tiap model — menganggur. Builder ini mengeluarkannya jadi
PNG yang bisa ditampilkan asisten INLINE lewat tool `diagram_wiring`
(kanal `exploded_images`, pola `wiring_ref`).

Tiga cara ekstrak (lib yang SUDAH ada — tanpa pymupdf):
  1. PDF raster embedded : pypdfium2 objek gambar (tipe 3) → PIL → PNG.
     Filter min(w,h)>=150 & area>=200*200; buang logo berulang (hash sama
     muncul >=3 halaman) via dedup SHA1.
  2. PDF skema VEKTOR     : abs pneumatik & listrik HOHAN nyaris tak punya
     raster → render HALAMAN penuh pypdfium2 @ ~200 DPI → PNG.
  3. Excel                : openpyxl `ws._images[*]._data()` → PNG per sheet
     (label = nama model sheet); dedup hash antar-sheet bilingual.

`index.json` record: {file, label, deskripsi, sinonim:[], sumber, halaman, tipe,
dicari, seed_teks}. `label/deskripsi/sinonim` di sini masih SEED otomatis —
dikurasi belakangan oleh pass AI-vision (baca PNG, tulis label+sinonim Indonesia).
`tipe` ∈ diagram|skema|pinout|foto_unit.

`dicari` = boleh muncul dari pencarian bebas tool `diagram_wiring`. Default:
pinout/skema/foto_unit = True (jumlah kecil, nilai tinggi, dikurasi vision);
diagram naratif (boschcn/tft, ratusan) = False → TIDAK membanjiri pencarian,
tapi tetap ditautkan per-halaman ke `manual_teks` (ditampilkan menyertai
jawaban `cari_manual`). Kurator vision boleh menaikkan diagram penting → True.

Deterministik & aman di-rerun: nama file `{slug}_p{hal:03d}_{idx}.png`,
index diurut by file, PNG dioptimasi Pillow. Deploy: `data/manual_media/`
di-scp (pola `data/wiring`), cache index per-mtime.

Jalankan dari root repo:  python backend/tools/build_manual_media.py
                          python backend/tools/build_manual_media.py --stats
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from pathlib import Path

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        _s.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[2]
SRC_DIR = ROOT / "data" / "manuals"
OUT_DIR = ROOT / "data" / "manual_media"

_MIN_SIDE = 150          # sisi terpendek minimal (px)
_MIN_AREA = 200 * 200    # luas minimal (px) — buang ikon/garis kecil
_LOGO_MIN_PAGES = 3      # hash sama muncul di >= N halaman → logo/dekorasi, buang
_RENDER_DPI = 200        # DPI render halaman vektor
_MAX_SIDE = 1800         # sisi terpanjang maksimal (px) — hemat ukuran, tetap terbaca

# ── Sumber PDF raster (ekstrak gambar embedded) — (potongan nama, slug, label, tipe) ──
_PDF_RASTER = (
    ("MC NATIONAL V ENGINE ECU PIN DEFINITION", "mcnat5",
     "Skema/pinout ECU Bosch mesin MC National V", "pinout"),
    ("【英文版】NBCU PIN DEFINITION", "nbcu",
     "Skema/pinout NBCU (body control unit)", "pinout"),
    ("NANOBCU PIN DEFINITION", "nanobcu",
     "Skema/pinout NanoBCU (body control unit)", "pinout"),
    ("ZF-AMT PIN DEFINITION", "zfamt",
     "Skema/pinout ZF-AMT (transmisi otomatis)", "pinout"),
    ("Manual_Sinotruk_MC_BOSCHECU_DH_CHINESE", "boschcn",
     "Diagram manual Bosch ECU mesin MC (Sinotruk)", "diagram"),
    # TFT NanoBCU (Indonesia) ada di subfolder skema/ — di-handle terpisah di main()
)

# ── Sumber PDF vektor (render halaman) — (path relatif, slug, label, tipe) ──
_PDF_VECTOR = (
    ("skema/abs_kb_6x4_traktor.pdf", "abs_kb",
     "Skema pneumatik ABS 6x4 traktor (tipe KB)", "skema"),
    ("skema/abs_wabco_6x4_traktor.pdf", "abs_wabco",
     "Skema pneumatik ABS WABCO 6x4 traktor", "skema"),
    ("skema/abs_6x4_nontraktor.pdf", "abs_rigid",
     "Skema pneumatik ABS 6x4 non-traktor (rigid/cargo)", "skema"),
    ("skema/listrik_hohan_howo_n_mc_a12.pdf", "listrik_hohan",
     "Skema kelistrikan HOHAN N / HOWO N mesin MC (A12)", "skema"),
)

# TFT NanoBCU: manual Indonesia (72 hal, 89 gambar layar/panel) — raster.
_TFT = ("skema/manual_tft_nanobcu.pdf", "tft",
        "Manual servis panel instrumen TFT NanoBCU (Indonesia)", "diagram")

_IMAGE_OBJ = 3  # pypdfium2: FPDF_PAGEOBJ_IMAGE
_DICARI_TIPE = {"pinout", "skema", "foto_unit"}  # boleh dicari bebas diagram_wiring


def _slugify(s: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "_", (s or "").strip()).strip("_").lower()
    return s or "x"


def _find(needle: str) -> Path | None:
    for p in sorted(SRC_DIR.glob("*.pdf")):
        if needle.lower() in p.name.lower():
            return p
    return None


def _page_text_snip(path: Path, npages: int) -> dict[int, str]:
    """Peta halaman(1-based) → cuplikan teks (bahan seed label untuk kurator)."""
    out: dict[int, str] = {}
    try:
        import pdfplumber
        with pdfplumber.open(path) as pl:
            for i, pg in enumerate(pl.pages):
                if i >= npages:
                    break
                t = (pg.extract_text(x_tolerance=1.5) or "").strip()
                if t:
                    out[i + 1] = " ".join(t.split())[:200]
    except Exception:
        pass
    return out


def _optimize_png(pil) -> bytes:
    """PIL → PNG bytes (RGB, dioptimasi) — deterministik (tanpa chunk waktu)."""
    from PIL import Image
    if pil.mode in ("RGBA", "P", "LA"):
        bg = Image.new("RGB", pil.size, (255, 255, 255))
        try:
            bg.paste(pil, mask=pil.convert("RGBA").split()[-1])
            pil = bg
        except Exception:
            pil = pil.convert("RGB")
    elif pil.mode != "RGB":
        pil = pil.convert("RGB")
    w, h = pil.size
    if max(w, h) > _MAX_SIDE:  # perkecil proporsional — hemat, tetap terbaca
        r = _MAX_SIDE / max(w, h)
        pil = pil.resize((max(1, int(w * r)), max(1, int(h * r))),
                         Image.LANCZOS)
    buf = io.BytesIO()
    pil.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def extract_pdf_raster(path: Path, slug: str, label: str, tipe: str) -> list[dict]:
    """Objek gambar embedded tiap halaman → kandidat {slug,halaman,idx,png,seed}."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        npages = len(pdf)
        snips = _page_text_snip(path, npages)
        # 1) kumpulkan semua kandidat + hash (untuk deteksi logo berulang)
        cand: list[dict] = []
        hash_pages: dict[str, set[int]] = {}
        for i in range(npages):
            page = pdf[i]
            try:
                objs = [o for o in page.get_objects(max_depth=4)
                        if o.type == _IMAGE_OBJ]
            except Exception:
                objs = []

            def _key(o):
                try:
                    l, b, r, t = o.get_pos()
                    return (-t, l)
                except Exception:
                    return (0.0, 0.0)

            objs.sort(key=_key)
            for idx, o in enumerate(objs):
                pil = None
                for render in (False, True):
                    try:
                        pil = o.get_bitmap(render=render).to_pil()
                        break
                    except Exception:
                        pil = None
                if pil is None:
                    continue
                w, h = pil.size
                if min(w, h) < _MIN_SIDE or w * h < _MIN_AREA:
                    continue
                png = _optimize_png(pil)
                hsh = hashlib.sha1(png).hexdigest()
                hash_pages.setdefault(hsh, set()).add(i + 1)
                cand.append({"halaman": i + 1, "idx": idx, "png": png,
                             "hash": hsh, "seed": snips.get(i + 1, "")})
    finally:
        pdf.close()

    # 2) buang logo berulang; simpan 1 file per hash unik (kemunculan pertama)
    seen: set[str] = set()
    rows: list[dict] = []
    for c in cand:
        if len(hash_pages.get(c["hash"], ())) >= _LOGO_MIN_PAGES:
            continue  # dekorasi/header/footer
        if c["hash"] in seen:
            continue  # duplikat identik antar-halaman
        seen.add(c["hash"])
        fname = f"{slug}_p{c['halaman']:03d}_{c['idx']}.png"
        rows.append({
            "file": fname, "png": c["png"],
            "label": f"{label} — hal {c['halaman']}",
            "deskripsi": "", "sinonim": [],
            "sumber": path.name, "halaman": c["halaman"],
            "tipe": tipe, "dicari": tipe in _DICARI_TIPE,
            "seed_teks": c["seed"],
        })
    return rows


def render_pdf_pages(path: Path, slug: str, label: str, tipe: str) -> list[dict]:
    """Render tiap halaman PDF vektor jadi PNG @ _RENDER_DPI (skema tanpa raster)."""
    import pypdfium2 as pdfium

    scale = _RENDER_DPI / 72.0
    pdf = pdfium.PdfDocument(str(path))
    rows: list[dict] = []
    try:
        npages = len(pdf)
        snips = _page_text_snip(path, npages)
        for i in range(npages):
            try:
                pil = pdf[i].render(scale=scale).to_pil()
            except Exception:
                continue
            png = _optimize_png(pil)
            suffix = f" — hal {i + 1}" if npages > 1 else ""
            rows.append({
                "file": f"{slug}_p{i + 1:03d}.png", "png": png,
                "label": f"{label}{suffix}",
                "deskripsi": "", "sinonim": [],
                "sumber": path.name, "halaman": i + 1,
                "tipe": tipe, "dicari": tipe in _DICARI_TIPE,
                "seed_teks": snips.get(i + 1, ""),
            })
    finally:
        pdf.close()
    return rows


def extract_excel(path: Path) -> list[dict]:
    """Foto unit embedded tiap sheet Excel → PNG (label = model sheet). Dedup
    hash antar-sheet (versi bilingual sering ulang foto sama)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    xslug = _slugify(path.stem)[:16]
    seen: set[str] = set()
    rows: list[dict] = []
    for ws in wb.worksheets:
        imgs = list(getattr(ws, "_images", []) or [])
        model = " ".join(str(ws.title).split())
        mslug = _slugify(ws.title)[:20]
        for k, im in enumerate(imgs):
            try:
                data = im._data()
            except Exception:
                continue
            if not data:
                continue
            try:
                from PIL import Image
                pil = Image.open(io.BytesIO(data))
                pil.load()
            except Exception:
                continue
            w, h = pil.size
            if min(w, h) < _MIN_SIDE or w * h < _MIN_AREA:
                continue
            png = _optimize_png(pil)
            hsh = hashlib.sha1(png).hexdigest()
            if hsh in seen:
                continue
            seen.add(hsh)
            rows.append({
                "file": f"xl_{xslug}_{mslug}_{k}.png", "png": png,
                "label": f"Foto unit {model}",
                "deskripsi": "", "sinonim": [],
                "sumber": path.name, "halaman": None,
                "tipe": "foto_unit", "dicari": True, "seed_teks": model,
            })
    return rows


def gather() -> list[dict]:
    rows: list[dict] = []
    # PDF raster
    for needle, slug, label, tipe in _PDF_RASTER:
        p = _find(needle)
        if p is None:
            print(f"⚠️ lewati (tak ada): {needle}")
            continue
        got = extract_pdf_raster(p, slug, label, tipe)
        print(f"  raster {slug:9s} {len(got):3d} gambar  ({p.name[:40]})")
        rows.extend(got)
    # TFT (subfolder skema/)
    ptft = SRC_DIR / _TFT[0]
    if ptft.exists():
        got = extract_pdf_raster(ptft, _TFT[1], _TFT[2], _TFT[3])
        print(f"  raster {_TFT[1]:9s} {len(got):3d} gambar  (TFT NanoBCU)")
        rows.extend(got)
    # PDF vektor (render halaman)
    for rel, slug, label, tipe in _PDF_VECTOR:
        p = SRC_DIR / rel
        if not p.exists():
            print(f"⚠️ lewati (tak ada): {rel}")
            continue
        got = render_pdf_pages(p, slug, label, tipe)
        print(f"  vektor {slug:9s} {len(got):3d} halaman ({p.name[:40]})")
        rows.extend(got)
    # Excel
    for p in sorted(SRC_DIR.glob("*.xlsx")):
        got = extract_excel(p)
        if got:
            print(f"  excel  {_slugify(p.stem)[:9]:9s} {len(got):3d} foto    ({p.name[:40]})")
        rows.extend(got)
    return rows


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--stats", action="store_true",
                    help="hitung saja, jangan tulis file")
    args = ap.parse_args()

    if not SRC_DIR.is_dir():
        raise SystemExit(f"⛔ dir sumber tak ada: {SRC_DIR}")
    rows = gather()
    if len(rows) < 50:
        raise SystemExit(f"⛔ hanya {len(rows)} gambar terekstrak — cek layout/lib?")

    # gabung dengan kurasi yang sudah ada (label/deskripsi/sinonim hasil pass
    # vision) — jangan timpa; cocok by nama file.
    idx_path = OUT_DIR / "index.json"
    kur: dict[str, dict] = {}
    if idx_path.exists():
        try:
            for r in json.loads(idx_path.read_text(encoding="utf-8")) or []:
                kur[r.get("file")] = r
        except Exception:
            kur = {}

    by_tipe: dict[str, int] = {}
    for r in rows:
        by_tipe[r["tipe"]] = by_tipe.get(r["tipe"], 0) + 1
    print(f"\nTotal {len(rows)} gambar; per tipe: {by_tipe}")

    if args.stats:
        return 0

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    # tulis PNG + rakit index (kurasi lama dipertahankan bila ada)
    index: list[dict] = []
    kept_files: set[str] = set()
    for r in rows:
        (OUT_DIR / r["file"]).write_bytes(r.pop("png"))
        kept_files.add(r["file"])
        prev = kur.get(r["file"])
        if prev:  # pertahankan kurasi manusia/vision
            for f in ("label", "deskripsi", "sinonim", "tipe"):
                if prev.get(f) not in (None, "", []):
                    r[f] = prev[f]
            if "dicari" in prev:  # kurator boleh promote/demote
                r["dicari"] = prev["dicari"]
        index.append(r)

    # buang PNG yatim (sumber berubah — file lama tak lagi diproduksi)
    for old in OUT_DIR.glob("*.png"):
        if old.name not in kept_files:
            old.unlink()

    index.sort(key=lambda r: r["file"])
    idx_path.write_text(
        json.dumps(index, ensure_ascii=False, indent=1), encoding="utf-8")
    total_kb = sum((OUT_DIR / r["file"]).stat().st_size for r in index) // 1024
    curated = sum(1 for r in index if r.get("sinonim"))
    print(f"✅ {idx_path.relative_to(ROOT)}: {len(index)} gambar "
          f"({total_kb} KB); terkurasi(sinonim): {curated}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
